#!/usr/bin/env python3
"""
sheet_to_docx.py
=================

Turn a Google Sheet of candidates (Name, ID, link to a scanned ID on
Google Drive) into a document with one page per candidate, showing their
Name, ID, and the scanned ID image. ID scans may be uploaded as JPG or
PDF — both are handled automatically. Output can be a Word (.docx) file,
a PDF, or both — you'll be asked which when you run the script (or pass
--format to skip the prompt).

------------------------------------------------------------------------
SETUP (one-time)
------------------------------------------------------------------------
1. Install dependencies:

       pip install pandas python-docx requests pillow pymupdf reportlab

   (pymupdf renders PDF ID scans to an image; reportlab builds the PDF
   output option. Both are required since your scans come in as JPG or
   PDF and you want the option of PDF output.)

2. Sharing permissions matter. This script does NOT use the Google API
   or OAuth — it reads your sheet and downloads images using plain HTTP,
   the same way a browser would. That means:

     - Your Google Sheet must be shared as "Anyone with the link -> Viewer"
       (File > Share > General access).
     - Every ID-scan file (or the whole Drive folder they're in) must
       also be shared as "Anyone with the link -> Viewer".

   If your data is sensitive and you don't want it publicly-by-link
   accessible, see the "PRIVATE SHEETS / DRIVE FILES" note near the
   bottom of this file for how to swap in an authenticated approach
   instead.

3. Find your Sheet ID. It's the long string in the sheet's URL:

       https://docs.google.com/spreadsheets/d/<THIS_PART_IS_THE_ID>/edit#gid=0

   If your data is on a specific tab (not the first one), also grab the
   "gid=" number from the end of the URL.

4. Edit the CONFIG block below: sheet_id, the list of columns you want
   printed on each page (text_columns), the link column, and the output
   name. Or pass them as command-line arguments — see USAGE.

------------------------------------------------------------------------
USAGE
------------------------------------------------------------------------
    python sheet_to_docx.py
        -> prompts you to choose Word / PDF / Both

    python sheet_to_docx.py --format pdf
    python sheet_to_docx.py --format both --sheet-id 1AbCDefGhIJKlmNoPQRstuVWxyz
    python sheet_to_docx.py --output candidates_2026 --format docx

    # Override which columns get printed (first one becomes the page title):
    python sheet_to_docx.py --column Name --column ID --column Department

Run with --help to see all options.
"""

import argparse
import io
import os
import re
import sys
import time

import pandas as pd
import requests
from PIL import Image, UnidentifiedImageError
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# ==========================================================================
# CONFIG — edit these defaults to match your sheet, or override via CLI
# ==========================================================================
CONFIG = {
    # The long ID from your Google Sheet's URL (between /d/ and /edit)
    "sheet_id": "1ffkoSIuMEzn1e2ZD24oosEII3TOQ76RFV0bzNi31HHI",

    # Optional: the gid= number if your data isn't on the first tab.
    # Leave as None to use the first/default tab.
    "gid": 1418612790,

    # Every column (exact header text from row 1 of your sheet) you want
    # printed on each candidate's page, in the order they should appear.
    # Add or remove as many as you like — it's not limited to Name/ID.
    #
    # The FIRST entry is shown as the big page title; every entry after
    # that is shown as a "Label: value" line beneath it.
    #
    # Each entry can just be the column header (used as its own label):
    #     "Name",
    # ...or a dict if you want the printed label to differ from the
    # sheet's header text:
    #     {"header": "Roll No", "label": "Roll Number"},
    "text_columns": [
        "Name as per Bank Account",
        "Name of Bank",
        "Bank Account Number",
        "Bank IFSC code",
        "PAN details",
    ],

    # The column holding the Google Drive link to the ID scan.
    "link_column": "Upload clear copy of PAN card in jpg or PDF format",

    # Base name (no extension) for the generated file(s). If --format is
    # "both", you'll get <output_base>.docx AND <output_base>.pdf.
    "output_base": "candidates",

    # "docx", "pdf", "both", or None to be asked at the terminal each run.
    "output_format": None,

    # How wide (in inches) each ID image should appear on the page.
    "image_width_inches": 4.5,

    # Where downloaded scans are cached before being embedded.
    "temp_dir": "temp_id_scans",
}


# ==========================================================================
# Column configuration
# ==========================================================================
def normalize_columns(column_defs: list) -> list[tuple[str, str]]:
    """Turn the text_columns config (a mix of plain header strings and
    {"header": ..., "label": ...} dicts) into a uniform list of
    (header, label) tuples, in the given order."""
    normalized = []
    for entry in column_defs:
        if isinstance(entry, dict):
            header = entry["header"]
            label = entry.get("label", header)
        else:
            header = entry
            label = entry
        normalized.append((header, label))
    return normalized


def clean_cell(value) -> str:
    """Turn a pandas cell value into a plain string, treating NaN/None
    (blank cells) as empty rather than the literal text 'nan'."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


# ==========================================================================
# Google Sheet loading
# ==========================================================================
def load_sheet(sheet_id: str, gid: str | None = None) -> pd.DataFrame:
    """Download a public Google Sheet as CSV and parse it into a DataFrame."""
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
    if gid:
        url += f"&gid={gid}"

    resp = requests.get(url, timeout=30)
    if resp.status_code != 200 or "text/html" in resp.headers.get("Content-Type", ""):
        raise RuntimeError(
            "Could not download the sheet as CSV. Most likely the sheet isn't "
            "shared as 'Anyone with the link' yet, or the sheet_id/gid is wrong.\n"
            f"Tried URL: {url}"
        )

    df = pd.read_csv(io.StringIO(resp.text), dtype=str)
    df = df.dropna(how="all")  # drop fully blank rows
    return df


# ==========================================================================
# Google Drive link handling
# ==========================================================================
DRIVE_ID_PATTERNS = [
    re.compile(r"/d/([a-zA-Z0-9_-]{10,})"),
    re.compile(r"[?&]id=([a-zA-Z0-9_-]{10,})"),
    re.compile(r"/open\?id=([a-zA-Z0-9_-]{10,})"),
]


def extract_drive_id(url: str) -> str | None:
    """Pull the file ID out of a Google Drive share link, in any common format."""
    if not isinstance(url, str) or not url.strip():
        return None
    for pattern in DRIVE_ID_PATTERNS:
        m = pattern.search(url)
        if m:
            return m.group(1)
    return None


def _get_confirm_token(response: requests.Response) -> str | None:
    """Google Drive shows a 'can't scan for viruses' interstitial for larger
    files. Find the confirm token so we can push through it."""
    for key, value in response.cookies.items():
        if key.startswith("download_warning"):
            return value
    content_type = response.headers.get("Content-Type", "")
    if "text/html" in content_type:
        m = re.search(r"confirm=([0-9A-Za-z_-]+)", response.text)
        if m:
            return m.group(1)
    return None


def download_drive_file(file_id: str, dest_path: str, session: requests.Session,
                         max_retries: int = 3) -> str:
    """Download a Google Drive file by ID to dest_path. Raises on failure."""
    url = "https://drive.google.com/uc?export=download"
    last_error = None

    for attempt in range(max_retries):
        try:
            resp = session.get(url, params={"id": file_id}, stream=True, timeout=30)
            token = _get_confirm_token(resp)
            if token:
                resp = session.get(
                    url, params={"id": file_id, "confirm": token}, stream=True, timeout=30
                )

            content_type = resp.headers.get("Content-Type", "")
            if "text/html" in content_type:
                snippet = resp.text[:200].replace("\n", " ")
                raise ValueError(
                    "Drive returned a web page instead of the file — the file is "
                    f"probably not shared as 'Anyone with the link'. (Response: {snippet}...)"
                )

            with open(dest_path, "wb") as f:
                for chunk in resp.iter_content(32768):
                    if chunk:
                        f.write(chunk)
            return dest_path

        except Exception as e:
            last_error = e
            if attempt < max_retries - 1:
                time.sleep(1.5 * (attempt + 1))

    raise RuntimeError(f"Failed to download file {file_id} after {max_retries} attempts: {last_error}")


# ==========================================================================
# ID scan handling — accepts JPG (or any plain image) and PDF uploads
# ==========================================================================
def sniff_file_kind(path: str) -> str:
    """Look at the first few bytes to tell PDF from image from unknown,
    regardless of what extension (or lack of one) the download has."""
    with open(path, "rb") as f:
        header = f.read(8)
    if header.startswith(b"%PDF-"):
        return "pdf"
    if header[:3] == b"\xff\xd8\xff":
        return "jpeg"
    if header[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    return "unknown"


def normalize_image(input_path: str, output_path: str, max_dimension: int = 1600) -> str | None:
    """Re-save a JPG/PNG as a clean JPEG that both Word and the PDF builder
    can always embed. Returns output_path on success, None if it isn't a
    valid image."""
    try:
        img = Image.open(input_path)
        img.load()
    except (UnidentifiedImageError, OSError):
        return None

    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    img.thumbnail((max_dimension, max_dimension))
    img.save(output_path, "JPEG", quality=88)
    return output_path


def pdf_first_page_to_image(input_path: str, output_path: str) -> str | None:
    """Render page 1 of a PDF ID scan to a PNG. Requires pymupdf."""
    import pymupdf as fitz  # checked for availability at startup, see main()
    try:
        doc = fitz.open(input_path)
        page = doc.load_page(0)
        pix = page.get_pixmap(dpi=200)
        pix.save(output_path)
        return output_path
    except Exception:
        return None


def prepare_candidate_image(link: str, row_index: int, temp_dir: str,
                             session: requests.Session, pdf_support: bool) -> tuple[str | None, str | None]:
    """Download one candidate's ID scan (JPG or PDF) and normalize it to a
    single JPEG ready to embed. Returns (image_path, error_message) —
    exactly one will be None."""
    file_id = extract_drive_id(link)
    if not file_id:
        return None, f"Could not parse a Google Drive file ID from link: {link!r}"

    raw_path = os.path.join(temp_dir, f"row{row_index}_raw")
    try:
        download_drive_file(file_id, raw_path, session)
    except Exception as e:
        return None, str(e)

    normalized_path = os.path.join(temp_dir, f"row{row_index}.jpg")
    kind = sniff_file_kind(raw_path)

    if kind == "pdf":
        if not pdf_support:
            return None, "Scan is a PDF but pymupdf isn't installed — run `pip install pymupdf`."
        page_image = os.path.join(temp_dir, f"row{row_index}_pdf.png")
        rendered = pdf_first_page_to_image(raw_path, page_image)
        if not rendered:
            return None, "Could not render the PDF's first page (file may be corrupt or password-protected)."
        result = normalize_image(rendered, normalized_path)
        if result:
            return result, None
        return None, "Rendered the PDF page but couldn't process the resulting image."

    # kind is "jpeg", "png", or "unknown" — just try opening it as an image.
    result = normalize_image(raw_path, normalized_path)
    if result:
        return result, None

    return None, "Downloaded file is neither a recognizable image (JPG/PNG) nor a PDF."


# ==========================================================================
# Word (.docx) document assembly
# ==========================================================================
def build_docx(records: list[dict], output_path: str, image_width_inches: float):
    doc = Document()

    for i, rec in enumerate(records):
        fields = rec.get("fields", [])
        image_path = rec.get("image_path")
        error = rec.get("error")

        if fields:
            title_label, title_value = fields[0]
            title_text = title_value or f"(Missing {title_label})"
            heading = doc.add_heading(str(title_text), level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for label, value in fields[1:]:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(f"{label}: {value or '(missing)'}")
                run.font.size = Pt(14)
        else:
            doc.add_heading("(No columns configured)", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()  # spacer

        if image_path and os.path.exists(image_path):
            img_para = doc.add_paragraph()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_para.add_run().add_picture(image_path, width=Inches(image_width_inches))
        else:
            err_para = doc.add_paragraph()
            err_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = err_para.add_run(f"[ID image not available: {error or 'unknown error'}]")
            run.italic = True

        if i < len(records) - 1:
            doc.add_page_break()

    doc.save(output_path)


# ==========================================================================
# PDF document assembly (built directly, not via docx->pdf conversion, so
# it works on any OS without LibreOffice/MS Word installed)
# ==========================================================================
def build_pdf(records: list[dict], output_path: str, image_width_inches: float):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    page_width, page_height = letter
    c = canvas.Canvas(output_path, pagesize=letter)

    for rec in records:
        fields = rec.get("fields", [])
        image_path = rec.get("image_path")
        error = rec.get("error")

        y = page_height - 1.2 * inch

        if fields:
            title_label, title_value = fields[0]
            title_text = str(title_value or f"(Missing {title_label})")
            c.setFont("Helvetica-Bold", 20)
            c.drawCentredString(page_width / 2, y, title_text)
            y -= 0.4 * inch

            c.setFont("Helvetica", 13)
            for label, value in fields[1:]:
                c.drawCentredString(page_width / 2, y, f"{label}: {value or '(missing)'}")
                y -= 0.3 * inch
        else:
            c.setFont("Helvetica-Bold", 16)
            c.drawCentredString(page_width / 2, y, "(No columns configured)")
            y -= 0.4 * inch

        y -= 0.3 * inch  # spacer before image
        image_top = y

        if image_path and os.path.exists(image_path):
            img = ImageReader(image_path)
            iw, ih = img.getSize()
            display_w = image_width_inches * inch
            display_h = display_w * ih / iw

            max_h = max(1.0 * inch, image_top - 0.5 * inch)  # leave a bottom margin
            if display_h > max_h:
                display_h = max_h
                display_w = display_h * iw / ih

            x = (page_width - display_w) / 2
            img_y = max(0.5 * inch, image_top - display_h)
            c.drawImage(img, x, img_y, width=display_w, height=display_h,
                        preserveAspectRatio=True, mask="auto")
        else:
            c.setFont("Helvetica-Oblique", 12)
            c.drawCentredString(page_width / 2, max(1.0 * inch, image_top - 1.0 * inch),
                                 f"[ID image not available: {error or 'unknown error'}]")

        c.showPage()

    c.save()


# ==========================================================================
# Terminal prompt for output format
# ==========================================================================
def prompt_for_format() -> str:
    while True:
        choice = input(
            "\nWhich output do you want?\n"
            "  [1] Word (.docx)\n"
            "  [2] PDF\n"
            "  [3] Both\n"
            "Enter 1, 2, or 3 (default 1): "
        ).strip()
        if choice in ("", "1"):
            return "docx"
        if choice == "2":
            return "pdf"
        if choice == "3":
            return "both"
        print("Please enter 1, 2, or 3.")


# ==========================================================================
# Main
# ==========================================================================
def main():
    parser = argparse.ArgumentParser(description="Build a one-page-per-candidate Word/PDF doc from a Google Sheet.")
    parser.add_argument("--sheet-id", default=CONFIG["sheet_id"])
    parser.add_argument("--gid", default=CONFIG["gid"])
    parser.add_argument("--column", action="append", dest="columns",
                         help="A column header to print on each page (repeatable, in order — "
                              "first one given becomes the page title). Passing this at all "
                              "overrides the CONFIG['text_columns'] list.")
    parser.add_argument("--link-column", default=CONFIG["link_column"])
    parser.add_argument("--output", default=CONFIG["output_base"],
                         help="Base name for the output file(s), without extension.")
    parser.add_argument("--format", choices=["docx", "pdf", "both"], default=CONFIG["output_format"],
                         help="Skip the terminal prompt and pick the output format directly.")
    parser.add_argument("--image-width", type=float, default=CONFIG["image_width_inches"])
    parser.add_argument("--temp-dir", default=CONFIG["temp_dir"])
    args = parser.parse_args()

    if args.sheet_id == "PASTE_YOUR_GOOGLE_SHEET_ID_HERE":
        sys.exit(
            "Set your sheet_id first — either edit CONFIG['sheet_id'] at the top "
            "of this file, or run with --sheet-id YOUR_SHEET_ID."
        )

    output_format = args.format or prompt_for_format()

    # Check dependencies up front so a missing package fails fast, not
    # halfway through a long run.
    pdf_support = True
    try:
        import pymupdf  # noqa: F401  (for PDF-format ID scans)
    except ImportError:
        pdf_support = False
        print("Note: pymupdf isn't installed, so any ID scans uploaded as PDF will be "
              "skipped (with a note in their place). Run `pip install pymupdf` to support them.")

    if output_format in ("pdf", "both"):
        try:
            import reportlab  # noqa: F401
        except ImportError:
            sys.exit("PDF output needs the `reportlab` package — run `pip install reportlab` and try again.")

    output_base = os.path.splitext(args.output)[0] if args.output.lower().endswith((".docx", ".pdf")) else args.output

    column_defs = [{"header": c, "label": c} for c in args.columns] if args.columns else CONFIG["text_columns"]
    field_defs = normalize_columns(column_defs)
    if not field_defs:
        sys.exit("No columns configured — add at least one to CONFIG['text_columns'] (or pass --column).")

    os.makedirs(args.temp_dir, exist_ok=True)

    print(f"\nLoading sheet {args.sheet_id} ...")
    df = load_sheet(args.sheet_id, args.gid)
    print(f"Loaded {len(df)} rows.")

    needed_columns = [header for header, _label in field_defs] + [args.link_column]
    missing = [c for c in needed_columns if c not in df.columns]
    if missing:
        sys.exit(
            f"Column(s) not found in the sheet: {missing}. Columns present: {list(df.columns)}"
        )

    session = requests.Session()
    records = []

    for idx, row in df.iterrows():
        fields = [(label, clean_cell(row.get(header, ""))) for header, label in field_defs]
        link = row.get(args.link_column, "")
        display_name = fields[0][1] if fields and fields[0][1] else f"row {idx + 1}"

        print(f"[{idx + 1}/{len(df)}] {display_name} ...", end=" ", flush=True)

        if not isinstance(link, str) or not link.strip():
            print("no link, skipping image")
            records.append({"fields": fields, "image_path": None, "error": "No link provided"})
            continue

        image_path, error = prepare_candidate_image(link, idx, args.temp_dir, session, pdf_support)
        if image_path:
            print("ok")
        else:
            print(f"FAILED ({error})")
        records.append({"fields": fields, "image_path": image_path, "error": error})

        time.sleep(0.4)  # be polite to Drive's servers

    outputs = []
    if output_format in ("docx", "both"):
        docx_path = f"{output_base}.docx"
        print(f"\nBuilding {docx_path} ...")
        build_docx(records, docx_path, args.image_width)
        outputs.append(docx_path)

    if output_format in ("pdf", "both"):
        pdf_path = f"{output_base}.pdf"
        print(f"Building {pdf_path} ...")
        build_pdf(records, pdf_path, args.image_width)
        outputs.append(pdf_path)

    ok = sum(1 for r in records if r["image_path"])
    print(f"\nDone. {ok}/{len(records)} images embedded successfully.")
    print("Saved: " + ", ".join(outputs))
    if ok < len(records):
        print(f"(Rows with issues were still added with a note in place of the image — "
              f"see the console log above, or check files in {args.temp_dir}/)")


if __name__ == "__main__":
    main()


# ==========================================================================
# PRIVATE SHEETS / DRIVE FILES
# ==========================================================================
# If you can't make the sheet or the ID scans public-by-link, replace the
# `load_sheet` and `download_drive_file` steps with an authenticated
# version using a Google Cloud service account:
#
#   pip install gspread google-auth
#
#   import gspread
#   gc = gspread.service_account(filename="service_account.json")
#   sheet = gc.open_by_key(SHEET_ID).sheet1
#   df = pd.DataFrame(sheet.get_all_records())
#
# ...and for Drive downloads, use the Drive API's files().get_media()
# with the same service account credentials (share the sheet and the
# Drive folder with the service account's email address, found inside
# service_account.json, as a Viewer).
